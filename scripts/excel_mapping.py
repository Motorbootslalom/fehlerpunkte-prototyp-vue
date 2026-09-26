#!/usr/bin/env -S uv run --script
# /// script
# requires-python = ">=3.11"
# dependencies = ["openpyxl>=3.1"]
# ///
"""Zuordnung Webseite → Excel-Auswertung, aus den Formeln der Arbeitsmappe abgeleitet.

Aufruf (uv lädt openpyxl selbst):
    uv run scripts/excel_mapping.py [pfad/zur/Auswertung.xlsm]
    npm run excel:mapping -- [pfad/zur/Auswertung.xlsm]
Ohne Pfad wird die einzige .xlsm in ../Auswertung genommen.

Ergebnis: src/config/excel-zuordnung.json - nur Blattnamen, Spalten und Zeilen,
keine Personendaten. Dazu ein Bericht mit Auffälligkeiten in der Mappe.

Grundlage ist je Klasse/Lauf das Blatt „EF K? nL W“: Es hat je Tor und Richtung
ein festes Spaltenpaar (1H, 2H, …, 1R) für die beiden Bojen-Seiten. Die Formeln
seiner Datenzeilen verraten, welche Vorlagen-Zelle (T1_3_5 …, T2_4_5 …, T5 …,
ST …, SCH …) hineinfließt. Die Seite einer Vorlagen-Spalte steht in ihrem
Zeile-8-Bezug auf Eingabe!L8…L11 (L8/L9 = Hin, L10/L11 = Rück) - bewusst nicht in
den Texten („H S“ …), die sich in „Eingabe“ jederzeit ändern lassen. Bojen, die
eine Vorlage nicht hat (z. B. Innenbojen Klasse E), stehen im EF-Blatt ohne
Formel: Dorthin wird direkt geschrieben.
"""
from __future__ import annotations

import collections
import datetime
import json
import re
import sys
import warnings
from pathlib import Path

import openpyxl
from openpyxl.utils import get_column_letter as L

warnings.filterwarnings('ignore')

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / 'src/config/excel-zuordnung.json'
KLASSEN = ['E', '1', '2', '3', '4', '5', '6', '7']
LAEUFE = [1, 2, 3]
EF_ZEILEN = 40  # Datenzeilen je EF-Blatt (7 … 46, siehe B1 = COUNTIF(C7:C46 …))

# Seiten: Eingabe!L8/L10 bzw. Ü!$Y$4 ist die eine, L9/L11 bzw. Ü!$Y$5 die andere
# Bojen-Seite. In der Webseite heißen sie seiteB/seiteA (Tor 1/3/5 beobachtet
# seiteA = die Seite des Start-Bojentors in T1_3_5 - wird unten geprüft).
SEITE = {'Y4': 'B', 'Y5': 'A'}
EINGABE_ZEILE = {8: ('H', 'Y4'), 9: ('H', 'Y5'), 10: ('R', 'Y4'), 11: ('R', 'Y5')}

# Rolle eines Blatts nach seinem Namensanfang (Klasse/Lauf/„ALC“ werden ersetzt).
ROLLEN = [
    (re.compile(r'^T1_3_5 '), 'tor135'),
    (re.compile(r'^T2_4_5 '), 'tor245'),
    (re.compile(r'^T5 '), 'tor5'),
    (re.compile(r'^SCH '), 'sch'),
    (re.compile(r'^ST K\S+ \d+L$'), 'steg'),
    (re.compile(r'^EZ '), 'zeit'),
]

REF = re.compile(r"(?:'([^']+)'|([A-Za-zÄÖÜäöüß_][\wÄÖÜäöüß]*))!\$?([A-Z]{1,3})\$?(\d+)")


def refs(f: str | None) -> list[tuple[str, str, int]]:
    return [((a or b), c, int(r)) for a, b, c, r in REF.findall(f or '')]


def formel(ws, zeile: int, spalte: int) -> str | None:
    v = ws.cell(zeile, spalte).value
    return v if isinstance(v, str) and v.startswith('=') else None


def rolle(blatt: str) -> str | None:
    return next((r for pat, r in ROLLEN if pat.search(blatt)), None)


def spalten_nr(buchstabe: str) -> int:
    n = 0
    for ch in buchstabe:
        n = n * 26 + ord(ch) - 64
    return n


def zeilen_bereiche(zeilen: list[int]) -> str:
    """[26, 27, 28, 31] → „26–28, 31“."""
    teile, start = [], None
    for i, z in enumerate(zeilen):
        if start is None:
            start = z
        if i + 1 == len(zeilen) or zeilen[i + 1] != z + 1:
            teile.append(str(start) if start == z else f'{start}–{z}')
            start = None
    return ', '.join(teile)


class Mappe:
    def __init__(self, pfad: Path):
        self.wb = openpyxl.load_workbook(pfad, read_only=False, keep_vba=False)
        self.auffaellig: list[str] = []

    def endziel(self, blatt: str, zelle: str) -> tuple[str, str]:
        """Folgt reinen Bezügen (=Eingabe!B2 → =Eingaben!C2) bis zur Wert-Zelle."""
        for _ in range(10):
            f = self.wb[blatt][zelle].value
            if not (isinstance(f, str) and REF.fullmatch(f.strip().lstrip('='))):
                break
            b, s, r = refs(f)[0]
            blatt, zelle = b, f'{s}{r}'
        return blatt, zelle

    # ---- EF-Blatt ---------------------------------------------------------
    def ef_spalten(self, name: str) -> list[dict]:
        ws = self.wb[name]
        out, gruppe, paar, anzahl = [], None, 0, collections.Counter()
        for c in range(8, ws.max_column + 1):
            kopf = ws.cell(5, c).value
            z6 = ws.cell(6, c).value
            if kopf is not None:
                code = re.sub(r'\s+', '', str(kopf))
                anzahl[code] += 1
                gruppe, paar = (code if anzahl[code] == 1 else f'{code}#{anzahl[code]}'), 0
            elif z6 is None:
                gruppe = None
            if gruppe is None:
                continue
            paar += 1
            f6 = formel(ws, 6, c) or ''
            # Seite aus Zeile 6 (Ü!$Y$4/$Y$5), sonst aus der Lage im Paar
            # (manche Blätter tragen dort nur Text wie „L“/„R“).
            seite = 'Y4' if '$Y$4' in f6 else 'Y5' if '$Y$5' in f6 else (
                ('Y4' if paar == 1 else 'Y5') if z6 is not None else None)
            muster, je_zeile = collections.Counter(), {}
            for r in range(7, 7 + EF_ZEILEN):
                f = formel(ws, r, c)
                key = tuple(sorted((b, s, rr - r) for b, s, rr in refs(f))) if f else ()
                je_zeile[r] = key
                muster[key] += 1
            haupt = muster.most_common(1)[0][0]
            loecher = [r for r, k in je_zeile.items() if k != haupt]
            if loecher and haupt:
                self.auffaellig.append(
                    f'{name}!{L(c)} ({gruppe}): Zeilen {zeilen_bereiche(loecher)} weichen von der Formel der übrigen ab')
            out.append(dict(spalte=L(c), gruppe=gruppe, seite=seite, links=haupt))
        return out

    def vorlagen_seite(self, blatt: str, spalte: str) -> tuple[str, str] | None:
        f = formel(self.wb[blatt], 8, spalten_nr(spalte)) or ''
        m = re.search(r'Eingabe!\$?L\$?(\d+)', f)
        return EINGABE_ZEILE.get(int(m.group(1))) if m else None

    # ---- eine Klasse / ein Lauf -------------------------------------------
    def zuordnung(self, k: str, lauf: int) -> dict:
        ef_name = f'EF K{k} {lauf}L W'
        ziele: dict[str, dict] = {}
        blaetter: dict[str, str] = {'ef': ef_name}

        def setze(schluessel: str, ziel: dict):
            if schluessel in ziele and ziele[schluessel] != ziel:
                self.auffaellig.append(f'{ef_name}: „{schluessel}“ doppelt ({ziele[schluessel]} / {ziel})')
            ziele[schluessel] = ziel

        for sp in self.ef_spalten(ef_name):
            g = sp['gruppe']
            if g in ('Platz', 'S-Nr.', 'Name', 'Vorname', 'Zeit', 'Punkte', 'ges.Punkte', 'Fehler-Schlüssel'):
                continue
            g = 'S1' if g == 'Spd' else g
            if g == 'Disq.':
                for b, s, _ in sp['links']:
                    if (ro := rolle(b)):
                        blaetter[ro] = b
                        setze(f'disq/{ro}', {'blatt': ro, 'spalte': s})
                continue
            if not sp['links']:
                # Keine Formel: Boje fehlt in der Vorlage → direkt ins EF-Blatt.
                if sp['seite']:
                    setze(f"{g}/{SEITE[sp['seite']]}", {'blatt': 'ef', 'spalte': sp['spalte']})
                continue
            for b, s, versatz in sp['links']:
                ro = rolle(b)
                if ro is None or versatz != 3:
                    self.auffaellig.append(f'{ef_name}!{sp["spalte"]}: unerwarteter Bezug {b}!{s} (Versatz {versatz})')
                    continue
                blaetter[ro] = b
                vs = self.vorlagen_seite(b, s)
                if sp['seite'] and vs and vs[1] != sp['seite']:
                    self.auffaellig.append(
                        f'{ef_name}!{sp["spalte"]}: Seite laut EF ({sp["seite"]}) ≠ laut {b}!{s}8 ({vs[1]})')
                seite = vs[1] if vs else sp['seite']
                punkte = {'StegAB': 'steg/ab', 'StegAN': 'steg/an', 'SCH': 'sch', '5': 'tor5'}
                schluessel = punkte.get(g) or (f'{g}/{SEITE[seite]}' if seite else g)
                setze(schluessel, {'blatt': ro, 'spalte': s})

        # Fehlercodes/Bemerkung der Vorlagen ohne Bojen (nach Spaltenkopf in Zeile 7;
        # „Bemerkung“ heißt in manchen Klassen „Bemerkungen“).
        kopf_ziele = {
            'steg': {'Fehler AB': 'steg/ab/fehler', 'Fehler AN': 'steg/an/fehler', 'Bemerkung': 'bemerkung/steg'},
            'tor5': {'Fehler': 'tor5/fehler', 'Bemerkung': 'bemerkung/tor5'},
            'sch': {'Fehler': 'sch/fehler', 'Bemerkung': 'bemerkung/sch'},
        }
        for ro, tabelle in kopf_ziele.items():
            if ro not in blaetter:
                continue
            ws = self.wb[blaetter[ro]]
            for c in range(3, ws.max_column + 1):
                kopf = str(ws.cell(7, c).value or '').strip()
                kopf = 'Bemerkung' if kopf.startswith('Bemerkung') else kopf
                if kopf in tabelle:
                    setze(tabelle[kopf], {'blatt': ro, 'spalte': L(c)})

        # Zeit: EF „Punkte“ = EZ!K7 = IF(I7="",86400*J7,I7) - I nimmt eine einzelne
        # Zeit als Punkte auf, J mittelt die Zeiten der drei Zeitnehmer (F/G/H).
        g_formel = formel(self.wb[ef_name], 7, 7) or ''
        ez_ref = next(((b, s) for b, s, _ in refs(g_formel) if rolle(b) == 'zeit'), None)
        if ez_ref:
            blaetter['zeit'] = ez_ref[0]
            wz = self.wb[ez_ref[0]]
            in_k = re.findall(r'\b([A-Z]{1,3})7\b', formel(wz, 7, spalten_nr(ez_ref[1])) or '')
            mittel = next((s for s in in_k if formel(wz, 7, spalten_nr(s))), None)  # J (Formel)
            punkte = next((s for s in in_k if s != mittel), None)  # I (Eingabe)
            in_j = re.findall(r'\b([A-Z]{1,3})7\b', formel(wz, 7, spalten_nr(mittel)) or '') if mittel else []
            zeiten = [s for s in dict.fromkeys(in_j) if not formel(wz, 7, spalten_nr(s))]
            if punkte and len(zeiten) == 3:
                setze('zeit', {'blatt': 'zeit', 'spalten': zeiten, 'punkte': punkte})
            else:
                self.auffaellig.append(f'{ez_ref[0]}: Zeitspalten nicht erkannt (Punkte={punkte}, Zeiten={zeiten})')

        return {'blaetter': blaetter, 'ziele': dict(sorted(ziele.items()))}

    # ---- Zeilen, Startnummern, Titel ----------------------------------------
    def zeilen(self, blaetter: dict[str, str]) -> dict[str, dict]:
        """Erste Datenzeile + Startnummern-Spalte je Rolle (Suche per Startnummer)."""
        out = {}
        for ro, name in blaetter.items():
            ws = self.wb[name]
            if ro in ('ef', 'zeit'):
                # Startnummer per Formel aus der Startliste 'ST K?'!B7
                for c in range(1, 6):
                    if re.search(r"'ST K[^']*'!\$?B\$?7\b", formel(ws, 7, c) or ''):
                        out[ro] = {'ersteZeile': 7, 'nummer': L(c)}
                        break
            else:
                for r in range(8, 16):
                    if re.search(r'Eingabe!', formel(ws, r, 2) or '') and r >= 10:
                        out[ro] = {'ersteZeile': r, 'nummer': 'B'}
                        break
            if ro not in out:
                self.auffaellig.append(f'{name}: erste Datenzeile/Startnummern-Spalte nicht gefunden')
        return out

    def startnummern(self, k: str, tor135: str) -> dict | None:
        f = formel(self.wb[tor135], 10, 2) or ''
        r = refs(f)
        return {'blatt': r[0][0], 'spalte': r[0][1], 'ersteZeile': r[0][2]} if r else None


def main() -> None:
    if len(sys.argv) > 1:
        pfad = Path(sys.argv[1])
    else:
        kandidaten = sorted((ROOT.parent / 'Auswertung').glob('*.xlsm'))
        if len(kandidaten) != 1:
            sys.exit(f'Bitte den Pfad zur Auswertung angeben (gefunden: {[k.name for k in kandidaten]}).')
        pfad = kandidaten[0]
    print(f'  Lese {pfad.name} …')
    m = Mappe(pfad)

    klassen: dict[str, dict] = {}
    for k in KLASSEN:
        je_lauf = {}
        for lauf in LAEUFE:
            z = m.zuordnung(k, lauf)
            # Blattnamen lauf-neutral machen, damit gleiche Läufe zusammenfallen
            neutral = json.loads(json.dumps(z).replace(f' {lauf}L', ' {lauf}L'))
            je_lauf[lauf] = (z, neutral)
        erster = je_lauf[1][1]
        gleich = all(je_lauf[l][1] == erster for l in LAEUFE)
        if not gleich:
            m.auffaellig.append(f'Klasse {k}: Läufe unterschiedlich aufgebaut - Zuordnung je Lauf')
        z1 = je_lauf[1][0]
        eintrag = {
            'startnummern': m.startnummern(k, z1['blaetter']['tor135']),
            'zeilen': m.zeilen(z1['blaetter']),
        }
        if gleich:
            eintrag.update(erster)
        else:
            eintrag['laeufe'] = {str(l): je_lauf[l][1] for l in LAEUFE}
        klassen[k] = eintrag

    # Seiten-Annahme prüfen: Start-Boje in T1_3_5 = seiteA
    start_a = klassen['3']['ziele'].get('Start/A', {})
    if start_a.get('blatt') != 'tor135':
        m.auffaellig.append('Seiten-Annahme verletzt: Start/A liegt nicht im Tor-1/3/5-Blatt')

    t135 = klassen['3']['blaetter']['tor135'].replace('{lauf}', '1')
    titel_blatt, titel_zelle = m.endziel(t135, 'B3')
    daten = {
        '_hinweis': 'Erzeugt von scripts/excel_mapping.py - nicht von Hand ändern.',
        'quelle': pfad.name,
        'erzeugt': datetime.date.today().isoformat(),
        'titel': {'blatt': titel_blatt, 'zelle': titel_zelle},
        'seiten': {'A': 'Eingabe!L9/L11 bzw. Ü!$Y$5', 'B': 'Eingabe!L8/L10 bzw. Ü!$Y$4'},
        'klassen': klassen,
    }
    OUT.write_text(json.dumps(daten, ensure_ascii=False, indent=2) + '\n', encoding='utf-8')
    print(f'  ✔ {OUT.relative_to(ROOT)} geschrieben')

    # Bericht: je Schlüssel die Ziele über alle Klassen
    schluessel = sorted({s for e in klassen.values() for s in e.get('ziele', {})})
    print('\n  ' + 'Schlüssel'.ljust(18) + ''.join(f'K{k}'.ljust(12) for k in KLASSEN))
    for s in schluessel:
        zeile = '  ' + s.ljust(18)
        for k in KLASSEN:
            z = klassen[k].get('ziele', {}).get(s)
            txt = '—' if not z else f"{z['blatt']}!{z.get('spalte') or '/'.join(z['spalten'])}"
            zeile += txt[:11].ljust(12)
        print(zeile)
    if m.auffaellig:
        print('\n  Auffälligkeiten in der Mappe:')
        for a in dict.fromkeys(m.auffaellig):
            print('   ⚠', a)
    else:
        print('\n  ✔ keine Auffälligkeiten in der Mappe')


if __name__ == '__main__':
    main()
